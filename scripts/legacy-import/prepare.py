"""
One-off extraction step for the legacy commercial spreadsheets into clean,
reviewable JSON. Run this once, inspect the JSON if you like, then run
`node scripts/legacy-import/run.mjs` to write it into Supabase.

Usage:
    python scripts/legacy-import/prepare.py

Reads (paths are the originals the user provided; edit here if they move):
    - Comercial 2026 - Base de Clientes.csv        (875 historical clients)
    - Parceiros Comerciais e Embaixadores.xlsx      (partners/ambassadors/coupons)
    - Links de Pagamento comercial externo.xlsx     (~37 legacy Asaas links)

Writes:
    scripts/legacy-import/data/clients.json
    scripts/legacy-import/data/partners.json
    scripts/legacy-import/data/ambassadors.json
    scripts/legacy-import/data/legacy_links.json
"""
import csv
import json
import os
import re

import openpyxl

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
OUT_DIR = os.path.join(BASE_DIR, "data")
os.makedirs(OUT_DIR, exist_ok=True)

CLIENTS_CSV = r"c:\Users\otavi\Downloads\⚖️ Chat Jurídico _ Comercial 2026 - ⭐ Base de Clientes.csv"
PARTNERS_XLSX = r"c:\Users\otavi\Downloads\Chat Jurídico _ Parceiros Comerciais e Embaixadores.xlsx"
LINKS_XLSX = r"c:\Users\otavi\Downloads\Links de Pagamento comercial externo.xlsx"


def clean(value):
    if value is None:
        return None
    if isinstance(value, str):
        value = value.strip()
        return value or None
    return value


def parse_date_br(value):
    """DD/MM/YYYY -> YYYY-MM-DD"""
    value = clean(value)
    if not value:
        return None
    match = re.match(r"^(\d{1,2})/(\d{1,2})/(\d{4})$", value)
    if not match:
        return None
    day, month, year = match.groups()
    return f"{year}-{month.zfill(2)}-{day.zfill(2)}"


def parse_installments(value):
    value = clean(value)
    if not value:
        return None
    match = re.search(r"(\d+)\s*x", value, re.IGNORECASE)
    return int(match.group(1)) if match else None


def parse_number(value):
    value = clean(value)
    if value is None:
        return None
    try:
        return float(str(value).replace(",", "."))
    except ValueError:
        return None


def parse_bool(value):
    value = clean(value)
    if value is None:
        return False
    return str(value).strip().upper() == "TRUE"


def extract_clients():
    with open(CLIENTS_CSV, encoding="utf-8-sig") as fh:
        rows = list(csv.DictReader(fh))

    out = []
    skipped_blank = 0
    for row in rows:
        status_raw = clean(row.get("STATUS"))
        if not status_raw:
            skipped_blank += 1
            continue
        status_map = {"Ativo": "ACTIVE", "Cancelado": "CANCELLED", "Congelada": "FROZEN"}
        status = status_map.get(status_raw)
        if not status:
            skipped_blank += 1
            continue

        period_raw = clean(row.get("Período"))
        billing_period = "ANNUAL" if period_raw == "ANUAL" else "MONTHLY"

        mrr_raw = clean(row.get("MRR"))
        try:
            mrr = float(str(mrr_raw).replace(",", ".")) if mrr_raw else 0.0
        except ValueError:
            mrr = 0.0
        subscription_value = round(mrr * 12, 2) if billing_period == "ANNUAL" else round(mrr, 2)

        out.append({
            "external_office_id": clean(row.get("office_id")),
            "office_name": clean(row.get("Nome do Escritório")) or "Sem nome",
            "responsible_name": clean(row.get("Nome do Responsável")),
            "email": clean(row.get("Email")),
            "phone": clean(row.get("Whatsapp Responsável")),
            "city": clean(row.get("Cidade")),
            "state": clean(row.get("Estado")),
            "service_area": clean(row.get("Área de Atendimento")),
            "status": status,
            "onboarding_completed": parse_bool(row.get("Onboarding")),
            "source_channel": clean(row.get("Veio de")) or clean(row.get("Parceiro")),
            "api_oficial": (clean(row.get("API Oficial")) or "").upper() == "SIM",
            "signed_at": parse_date_br(row.get("Assinado em")),
            "cancelled_at": parse_date_br(row.get("Cancelado em")),
            "cancellation_category": clean(row.get("Categoria do Cancelamento")),
            "cancellation_reason": clean(row.get("Motivo do Cancelamento")),
            "comments": clean(row.get("Comentários")),
            "implementation_date": parse_date_br(row.get("Data Implantacao")),
            "implementation_value": parse_number(row.get("Valor Implantacao")),
            "plan_name_raw": clean(row.get("Plano")),
            "billing_period": billing_period,
            "value": subscription_value,
            "payment_method": clean(row.get("Pagamento")),
            "installments": parse_installments(row.get("Parcelas")),
        })

    print(f"clients: {len(out)} usable rows, {skipped_blank} blank/filler rows skipped")
    with open(os.path.join(OUT_DIR, "clients.json"), "w", encoding="utf-8") as fh:
        json.dump(out, fh, ensure_ascii=False, indent=2)


def extract_partners_and_ambassadors():
    wb = openpyxl.load_workbook(PARTNERS_XLSX, data_only=True)

    def sheet_rows(sheet_name):
        ws = wb[sheet_name]
        header = [clean(cell.value) for cell in next(ws.iter_rows(min_row=1, max_row=1))]
        rows = []
        for raw in ws.iter_rows(min_row=2, values_only=True):
            record = {header[i]: clean(v) for i, v in enumerate(raw) if i < len(header) and header[i]}
            if any(record.values()):
                rows.append(record)
        return rows

    partners = []
    for row in sheet_rows("Parceiros"):
        name = row.get("Nome")
        if not name:
            continue
        partners.append({
            "name": name,
            "coupon": row.get("Cupom"),
            "reference_code": row.get("Ref"),
            "notes": row.get("Observações"),
        })

    ambassadors = []
    for row in sheet_rows("Embaixadores"):
        name = row.get("Nome")
        if not name:
            continue
        ambassadors.append({
            "name": name,
            "coupon": row.get("Cupom"),
            "reference_code": row.get("Ref"),
            "notes": row.get("Observações"),
        })

    print(f"partners: {len(partners)}, ambassadors: {len(ambassadors)}")
    with open(os.path.join(OUT_DIR, "partners.json"), "w", encoding="utf-8") as fh:
        json.dump(partners, fh, ensure_ascii=False, indent=2)
    with open(os.path.join(OUT_DIR, "ambassadors.json"), "w", encoding="utf-8") as fh:
        json.dump(ambassadors, fh, ensure_ascii=False, indent=2)


def extract_legacy_links():
    wb = openpyxl.load_workbook(LINKS_XLSX, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = []
    for raw in ws.iter_rows(min_row=2, values_only=True):
        if not raw or not clean(raw[0]):
            continue
        actor_name, plan_text, value, url, description = raw[0], raw[1], raw[2], raw[3], raw[4]
        actor_name, plan_text, description = clean(actor_name), clean(plan_text), clean(description)
        url = clean(url)
        billing_period = "ANNUAL" if plan_text and re.search(r"anual", plan_text, re.IGNORECASE) else "MONTHLY"
        installments = parse_installments(plan_text) if billing_period == "ANNUAL" else None
        asaas_id = None
        if url:
            match = re.search(r"/c/([A-Za-z0-9]+)", url)
            asaas_id = match.group(1) if match else None
        rows.append({
            "actor_name": actor_name,
            "plan_text": plan_text,
            "value": float(value) if value else None,
            "url": url,
            "asaas_payment_link_id": asaas_id,
            "display_name": description or plan_text,
            "billing_period": billing_period,
            "max_installments": installments,
        })

    print(f"legacy_links: {len(rows)}")
    with open(os.path.join(OUT_DIR, "legacy_links.json"), "w", encoding="utf-8") as fh:
        json.dump(rows, fh, ensure_ascii=False, indent=2)


if __name__ == "__main__":
    extract_clients()
    extract_partners_and_ambassadors()
    extract_legacy_links()
    print("Done. JSON written to scripts/legacy-import/data/")
